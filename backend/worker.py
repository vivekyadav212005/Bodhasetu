# backend/worker.py
import io
import fitz  # PyMuPDF
import openpyxl
import numpy as np
from datetime import datetime
from backend import db
from backend.utils.hash_utils import compute_sha256_bytes
from backend.extractors.ocr_tesseract import ocr_image_bytes
from backend.chunking.chunker import chunk_text_with_map as chunk_text
from backend.embeddings.embedder import embed_texts
from backend.qdrant_upsert import ensure_collection, upsert_chunks_vectors
from backend.utils.file_type import detect_basic_type


async def process_bytes(file_bytes: bytes, filename: str, department: str, source_meta: dict):
    """
    Full pipeline:
    1. Hash file -> dedup check
    2. Detect type -> extract text
    3. Chunk text -> save to Mongo
    4. Embed -> save to Qdrant
    """

    # ---------- 1. Compute hash & dedup check ----------
    sha256 = compute_sha256_bytes(file_bytes)
    existing = await db.get_document("documents", sha256)
    if existing:
        # Duplicate file: just update metadata
        await db.update_document(
            "documents",
            sha256,
            {"sources": existing.get("sources", []) + [source_meta],
             "last_seen": datetime.utcnow()}
        )
        return {"status": "duplicate", "doc_id": sha256}

    # ---------- 2. Save base doc metadata ----------
    doc_data = {
        "_id": sha256,
        "filename": filename,
        "department": department,
        "sources": [source_meta],
        "created_at": datetime.utcnow(),
        "updated_at": datetime.utcnow(),
    }
    await db.insert_document("documents", doc_data)

    # ---------- 3. Detect type ----------
    file_type = detect_basic_type(file_bytes, filename)
    extracted_text = ""
    chunks_meta = []

    if file_type == "pdf":
        extracted_text, chunks_meta = extract_pdf(file_bytes)
    elif file_type == "excel":
        extracted_text, chunks_meta = extract_excel(file_bytes)
    else:
        extracted_text = file_bytes.decode(errors="ignore")
        chunks_meta = [{"text": extracted_text, "chunk_index": 0}]

    # ---------- 4. Store chunks in Mongo ----------
    chunk_models = []
    for cm in chunks_meta:
        cm["_id"] = f"{sha256}_{cm['chunk_index']}"
        cm["doc_id"] = sha256
        await db.insert_document("chunks", cm)
        chunk_models.append(cm)

    # ---------- 5. Generate embeddings ----------
    texts = [cm["text"] for cm in chunk_models]
    if texts:
        vectors = embed_texts(texts)
        ensure_collection(vectors.shape[1])
        upsert_chunks_vectors(sha256, chunk_models, vectors)

    # ---------- 6. Final update ----------
    await db.update_document(
        "documents",
        sha256,
        {"updated_at": datetime.utcnow(), "num_chunks": len(chunk_models)}
    )

    return {"status": "new", "doc_id": sha256, "chunks": len(chunk_models)}


# ----------- Extraction Helpers -----------

def extract_pdf(file_bytes: bytes):
    """Extract text (and OCR for images) from PDFs"""
    doc = fitz.open(stream=file_bytes, filetype="pdf")
    chunks_meta = []
    texts = []
    idx = 0

    for page_num, page in enumerate(doc):
        text = page.get_text()
        if not text.strip():
            # Try OCR for scanned pages
            pix = page.get_pixmap()
            ocr_text = ocr_image_bytes(pix.tobytes("png"))
            text = ocr_text
        chunks = chunk_text(text, page_number=page_num)
        for c in chunks:
            c["chunk_index"] = idx
            idx += 1
            chunks_meta.append(c)
            texts.append(c["text"])
    return "\n".join(texts), chunks_meta


def extract_excel(file_bytes: bytes):
    """Extract text from Excel sheets"""
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    texts = []
    chunks_meta = []
    idx = 0

    for sheet in wb.sheetnames:
        ws = wb[sheet]
        sheet_texts = []
        for row in ws.iter_rows(values_only=True):
            row_text = " ".join([str(c) for c in row if c is not None])
            if row_text.strip():
                sheet_texts.append(row_text)
        sheet_text = "\n".join(sheet_texts)
        chunks = chunk_text(sheet_text, page_number=sheet)
        for c in chunks:
            c["chunk_index"] = idx
            idx += 1
            chunks_meta.append(c)
            texts.append(c["text"])

    return "\n".join(texts), chunks_meta